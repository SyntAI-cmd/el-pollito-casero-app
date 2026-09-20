"""Consolidado del día en Excel: una fila por pedido, importes como números con dos decimales."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.dinero import CERO
from app.modules.documentos.datos import Contexto

COLUMNAS = [
    ("N° remito", 11),
    ("Cliente", 32),
    ("CUIT", 15),
    ("Detalle", 48),
    ("Cajones", 9),
    ("Kilos", 10),
    ("Saldo cta. cte.", 15),
    ("Cajas adeudadas", 15),
    ("Total", 14),
    ("Preventista", 18),
    ("Turno", 9),
    ("Estado", 12),
]


def render(ctx: Contexto) -> bytes:
    libro = Workbook()
    hoja = libro.create_sheet("Consolidado", 0)
    assert isinstance(hoja, Worksheet)
    if libro.sheetnames[1:]:
        libro.remove(libro[libro.sheetnames[1]])
    fecha = ctx.fecha.strftime("%d/%m/%Y") if ctx.fecha else "varios días"
    hoja["A1"] = f"{ctx.fiscal.fiscal_razon_social} · Consolidado {fecha}"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append([])
    hoja.append([nombre for nombre, _ in COLUMNAS])
    for celda in hoja[3]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1A0706")
    for i, (_, ancho) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho

    for p in ctx.pedidos:
        datos = ctx.clientes.get(p.cliente_id)
        detalle = " · ".join(
            f"{i.producto_nombre} {i.cajas}c"
            if i.cajas
            else f"{i.producto_nombre} {i.kg_pedidos or ''}kg"
            for i in p.items
        )
        kilos = sum((i.kg_pesados or CERO for i in p.items), CERO)
        hoja.append(
            [
                p.numero,
                p.cliente_nombre,
                datos.cuit if datos else None,
                detalle,
                p.cajones,
                float(kilos),
                float(datos.saldo) if datos else None,
                datos.envases if datos else None,
                float(p.total),
                ctx.nombre(p.preventista_id),
                p.turno,
                p.estado,
            ]
        )
    ultima = hoja.max_row
    for fila in hoja.iter_rows(min_row=4, max_row=ultima):
        fila[5].number_format = "#,##0.000"
        fila[6].number_format = '"$" #,##0.00'
        fila[8].number_format = '"$" #,##0.00'
    if ctx.pedidos:
        hoja.append([])
        total = hoja.max_row + 1
        hoja.cell(row=total, column=1, value="TOTAL").font = Font(bold=True)
        hoja.cell(row=total, column=5, value=f"=SUM(E4:E{ultima})")
        hoja.cell(row=total, column=6, value=f"=SUM(F4:F{ultima})").number_format = "#,##0.000"
        hoja.cell(row=total, column=9, value=f"=SUM(I4:I{ultima})").number_format = '"$" #,##0.00'
    hoja.freeze_panes = "A4"
    for celda in hoja["D"]:
        celda.alignment = Alignment(wrap_text=True)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
