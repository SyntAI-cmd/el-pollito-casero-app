import os
import uuid
from io import BytesIO
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.integrations.storage import StorageLocal, usar_storage
from app.modules.auth.models import Usuario
from app.workers.cola import esperar_inline
from tests.conftest import HOY, pedido_base


@pytest.fixture(autouse=True)
def storage_temporal(tmp_path: Path) -> None:
    usar_storage(StorageLocal(tmp_path / "archivos"))


@pytest.fixture
async def dia_con_pedidos(
    cliente: AsyncClient,
    como_admin: dict[str, str],
    como_preventista: dict[str, str],
    cliente_con_precios: dict[str, str],
) -> list[dict[str, object]]:
    lista = []
    for n in range(5):
        pedido = (
            await cliente.post(
                "/pedidos",
                json=pedido_base(cliente_con_precios["id"], observaciones=f"Obs {n}"),
                headers=como_preventista if n % 2 else como_admin,
            )
        ).json()
        await cliente.post(
            f"/pedidos/{pedido['id']}/cajones/lote",
            json={
                "lote_id": str(uuid.uuid4()),
                "producto_codigo": "entero",
                "cajas": 3,
                "bruto_total": "65.1",
            },
            headers=como_admin,
        )
        lista.append(pedido)
    return lista


async def pedir(
    cliente: AsyncClient, headers: dict[str, str], **datos: object
) -> dict[str, object]:
    respuesta = await cliente.post("/documentos", json=datos, headers=headers)
    assert respuesta.status_code == 202, respuesta.text
    await esperar_inline()
    listo = await cliente.get(f"/documentos/{respuesta.json()['id']}", headers=headers)
    assert listo.json()["estado"] == "listo", listo.json()["error"]
    assert listo.json()["url"]
    return listo.json()


async def descargar(cliente: AsyncClient, url: str) -> bytes:
    respuesta = await cliente.get(url.replace("http://localhost:8000", ""))
    assert respuesta.status_code == 200
    return respuesta.content


@pytest.mark.parametrize("tipo", ["remitos", "hoja_pedidos", "hoja_ruta", "tickets"])
async def test_los_pdf_del_dia_se_generan_en_el_worker(
    cliente: AsyncClient,
    como_admin: dict[str, str],
    dia_con_pedidos: list[dict[str, object]],
    tipo: str,
) -> None:
    documento = await pedir(cliente, como_admin, tipo=tipo, fecha=HOY.isoformat())
    assert documento["nombre_archivo"] == f"{tipo}-{HOY.isoformat()}.pdf"
    contenido = await descargar(cliente, documento["url"])
    assert contenido.startswith(b"%PDF")
    paginas = contenido.count(b"/Type /Page\n") or contenido.count(b"/Type /Page")
    assert paginas >= 1


async def test_remitos_de_pedidos_puntuales_van_de_a_cuatro_por_hoja(
    cliente: AsyncClient, como_admin: dict[str, str], dia_con_pedidos: list[dict[str, object]]
) -> None:
    documento = await pedir(
        cliente, como_admin, tipo="remitos", pedido_ids=[p["id"] for p in dia_con_pedidos]
    )
    contenido = await descargar(cliente, documento["url"])
    # 5 remitos → 2 hojas A4.
    assert contenido.count(b"/Type /Page") - contenido.count(b"/Type /Pages") == 2


async def test_remitos_en_10x15_van_uno_por_pagina(
    cliente: AsyncClient,
    como_admin: dict[str, str],
    dia_con_pedidos: list[dict[str, object]],
) -> None:
    documento = await pedir(
        cliente,
        como_admin,
        tipo="remitos",
        formato="10x15",
        pedido_ids=[p["id"] for p in dia_con_pedidos],
    )
    contenido = await descargar(cliente, documento["url"])
    assert contenido.count(b"/Type /Page") - contenido.count(b"/Type /Pages") == 5
    assert b"/MediaBox [ 0 0 283.46" in contenido  # 100 × 150 mm en puntos
    if muestra := os.environ.get("MUESTRA_REMITOS"):  # para mirar el PDF a ojo
        a4 = await pedir(
            cliente, como_admin, tipo="remitos", pedido_ids=[p["id"] for p in dia_con_pedidos]
        )
        contenido_a4 = await descargar(cliente, a4["url"])
        Path(muestra).write_bytes(contenido)  # noqa: ASYNC240
        Path(muestra).with_name("remitos-a4.pdf").write_bytes(contenido_a4)  # noqa: ASYNC240


async def test_consolidado_excel_tiene_una_fila_por_pedido(
    cliente: AsyncClient, como_admin: dict[str, str], dia_con_pedidos: list[dict[str, object]]
) -> None:
    documento = await pedir(cliente, como_admin, tipo="consolidado", fecha=HOY.isoformat())
    libro = load_workbook(BytesIO(await descargar(cliente, documento["url"])))
    hoja = libro["Consolidado"]
    encabezados = [c.value for c in hoja[3]]
    assert encabezados[:9] == [
        "N° remito",
        "Cliente",
        "CUIT",
        "Detalle",
        "Cajones",
        "Kilos",
        "Saldo cta. cte.",
        "Cajas adeudadas",
        "Total",
    ]
    filas = [r for r in hoja.iter_rows(min_row=4, values_only=True) if r[0] and r[0] != "TOTAL"]
    assert len(filas) == 5
    assert filas[0][0] == dia_con_pedidos[0]["numero"]
    assert filas[0][4] == 3 and filas[0][5] == 60.0 and filas[0][8] == 300000.0


async def test_el_preventista_solo_imprime_lo_suyo_y_el_cobrador_nada(
    cliente: AsyncClient,
    como_preventista: dict[str, str],
    como_cobrador: dict[str, str],
    preventista: Usuario,
    dia_con_pedidos: list[dict[str, object]],
) -> None:
    ajeno = await cliente.post(
        "/documentos",
        json={"tipo": "hoja_ruta", "fecha": HOY.isoformat(), "preventista_id": str(uuid.uuid4())},
        headers=como_preventista,
    )
    assert ajeno.status_code == 403
    assert (
        await cliente.post(
            "/documentos", json={"tipo": "remitos", "fecha": HOY.isoformat()}, headers=como_cobrador
        )
    ).status_code == 403
    propio = await pedir(cliente, como_preventista, tipo="hoja_ruta", fecha=HOY.isoformat())
    assert propio["parametros"]["preventista_id"] == str(preventista.id)
    listado = await cliente.get(
        "/documentos", params={"fecha": HOY.isoformat()}, headers=como_preventista
    )
    assert [d["id"] for d in listado.json()] == [propio["id"]]
